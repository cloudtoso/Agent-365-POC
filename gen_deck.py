from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)

# === COWORK THEME COLOR PALETTE ===
NAVY = RGBColor(0x00, 0x20, 0x50)        # Dark Navy background
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_BLUE = RGBColor(0x00, 0x78, 0xD4)  # Primary accent
ACCENT_LTBLUE = RGBColor(0x50, 0xE6, 0xFF) # Secondary accent (light blue)
HIGHLIGHT_BG = RGBColor(0xE5, 0xF1, 0xFB)  # Content highlight box bg
BLACK = RGBColor(0x00, 0x00, 0x00)

# Microsoft 4-square logo colors
MS_RED = RGBColor(0xF2, 0x50, 0x22)
MS_GREEN = RGBColor(0x7F, 0xBA, 0x00)
MS_BLUE = RGBColor(0x00, 0xA4, 0xEF)
MS_YELLOW = RGBColor(0xFF, 0xB9, 0x00)

SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)

# Font families
FONT_LIGHT = 'Segoe UI Light'
FONT_SEMIBOLD = 'Segoe UI Semibold'
FONT_REGULAR = 'Segoe UI'

# === HELPERS ===
def no_line(shape):
    shape.line.fill.background()

def add_ms_logo(slide, left, top, sq_size):
    """Add Microsoft 4-square logo at given position. sq_size = size of each square in Pt."""
    gap = Pt(1.5)
    sz = sq_size
    colors = [MS_RED, MS_GREEN, MS_BLUE, MS_YELLOW]
    # Top-left red, Top-right green, Bottom-left blue, Bottom-right yellow
    positions = [
        (left, top),
        (left + sz + gap, top),
        (left, top + sz + gap),
        (left + sz + gap, top + sz + gap),
    ]
    for i, (l, t) in enumerate(positions):
        sq = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, sz, sz)
        sq.fill.solid()
        sq.fill.fore_color.rgb = colors[i]
        no_line(sq)

def add_footer(slide):
    """Cowork theme footer: tiny 4-square bottom-left, 'Microsoft' text, right-aligned identifier."""
    # Tiny 4-square logo (5x5pt squares)
    add_ms_logo(slide, Inches(0.4), Inches(7.0), Pt(5))
    # "Microsoft" text
    ms_box = slide.shapes.add_textbox(Inches(0.65), Inches(6.88), Inches(1.5), Inches(0.3))
    tf = ms_box.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = 'Microsoft'
    p.font.size = Pt(11)
    p.font.color.rgb = BLACK
    p.font.name = FONT_REGULAR
    p.alignment = PP_ALIGN.LEFT
    # Right-aligned identifier
    id_box = slide.shapes.add_textbox(Inches(8.5), Inches(6.88), Inches(4.5), Inches(0.3))
    tf2 = id_box.text_frame
    tf2.word_wrap = False
    p2 = tf2.paragraphs[0]
    p2.text = 'Agent 365  |  Proof of Concept'
    p2.font.size = Pt(10)
    p2.font.color.rgb = BLACK
    p2.font.name = FONT_REGULAR
    p2.alignment = PP_ALIGN.RIGHT

def add_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text

# === TITLE SLIDE ===
def make_title_slide(main_title, subtitle, tagline, customer_line, notes=''):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Dark navy background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = NAVY

    # Microsoft 4-square logo top-left (~19x19pt squares)
    add_ms_logo(slide, Inches(0.5), Inches(0.4), Pt(19))

    # Two horizontal bars top-right area
    # Blue bar ~230pt wide x 13pt tall
    blue_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(9.0), Inches(0.7), Pt(230), Pt(13)
    )
    blue_bar.fill.solid()
    blue_bar.fill.fore_color.rgb = ACCENT_BLUE
    no_line(blue_bar)

    # Light blue bar ~115pt wide x 13pt below it
    lt_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(9.0), Inches(0.92), Pt(115), Pt(13)
    )
    lt_bar.fill.solid()
    lt_bar.fill.fore_color.rgb = ACCENT_LTBLUE
    no_line(lt_bar)

    # Hero text 80pt Segoe UI Light white centered
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.0), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = main_title
    p.font.size = Pt(80)
    p.font.color.rgb = WHITE
    p.font.name = FONT_LIGHT
    p.font.bold = False
    p.alignment = PP_ALIGN.CENTER

    # Thin accent line (#50E6FF ~43x6pt) above the subtitle text
    accent_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(5.5), Inches(4.2), Pt(43), Pt(6)
    )
    accent_line.fill.solid()
    accent_line.fill.fore_color.rgb = ACCENT_LTBLUE
    no_line(accent_line)

    # Subtitle text
    sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(4.4), Inches(11.0), Inches(0.6))
    tf2 = sub_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(28)
    p2.font.color.rgb = WHITE
    p2.font.name = FONT_LIGHT
    p2.alignment = PP_ALIGN.CENTER

    # Tagline
    tag_box = slide.shapes.add_textbox(Inches(1.0), Inches(5.1), Inches(11.0), Inches(0.5))
    tf3 = tag_box.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    p3.text = tagline
    p3.font.size = Pt(14)
    p3.font.color.rgb = WHITE
    p3.font.name = FONT_REGULAR
    p3.alignment = PP_ALIGN.CENTER

    # Customer/date line
    cust_box = slide.shapes.add_textbox(Inches(1.0), Inches(5.6), Inches(11.0), Inches(0.4))
    tf4 = cust_box.text_frame
    tf4.word_wrap = True
    p4 = tf4.paragraphs[0]
    p4.text = customer_line
    p4.font.size = Pt(12)
    p4.font.color.rgb = ACCENT_LTBLUE
    p4.font.name = FONT_REGULAR
    p4.alignment = PP_ALIGN.CENTER

    if notes:
        add_notes(slide, notes)
    return slide

# === SECTION DIVIDER ===
section_counter = [0]

def section_divider(title, subtitle=''):
    section_counter[0] += 1
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Dark navy background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = NAVY

    # Full-height blue left bar (18pt wide, #0078D4)
    left_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Pt(18), SLIDE_H
    )
    left_bar.fill.solid()
    left_bar.fill.fore_color.rgb = ACCENT_BLUE
    no_line(left_bar)

    # Microsoft 4-square logo (13x13pt squares) top-right corner
    add_ms_logo(slide, Inches(12.0), Inches(0.3), Pt(13))

    # Large section number in light color (~380pt)
    num_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(6.0), Inches(5.5))
    tf = num_box.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = str(section_counter[0])
    p.font.size = Pt(380)
    p.font.color.rgb = RGBColor(0x00, 0x3A, 0x70)  # Slightly lighter navy for number
    p.font.name = FONT_LIGHT
    p.font.bold = False
    p.alignment = PP_ALIGN.LEFT

    # Section title in Segoe UI Light 56pt white
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(3.5), Inches(10.0), Inches(1.2))
    tf2 = title_box.text_frame
    tf2.word_wrap = True
    tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
    p2 = tf2.paragraphs[0]
    p2.text = title
    p2.font.size = Pt(56)
    p2.font.color.rgb = WHITE
    p2.font.name = FONT_LIGHT
    p2.font.bold = False
    p2.alignment = PP_ALIGN.LEFT

    if subtitle:
        p3 = tf2.add_paragraph()
        p3.text = subtitle
        p3.font.size = Pt(22)
        p3.font.color.rgb = ACCENT_LTBLUE
        p3.font.name = FONT_REGULAR
        p3.alignment = PP_ALIGN.LEFT
        p3.space_before = Pt(8)

    return slide

# === CONTENT SLIDE ===
def content_slide(title, bullets, notes=''):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # White background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = WHITE

    # Blue vertical accent bar left side (~8.6pt wide x ~40pt tall, #0078D4)
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Pt(8.6), Pt(40)
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = ACCENT_BLUE
    no_line(accent_bar)

    # Heading in Segoe UI Semibold 26pt black next to accent bar
    title_box = slide.shapes.add_textbox(Inches(0.72), Inches(0.40), Inches(11.5), Inches(0.7))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.color.rgb = BLACK
    p.font.name = FONT_SEMIBOLD
    p.font.bold = False
    p.alignment = PP_ALIGN.LEFT

    # Body content - bullets
    txBox = slide.shapes.add_textbox(Inches(0.72), Inches(1.4), Inches(11.5), Inches(5.0))
    tf2 = txBox.text_frame
    tf2.word_wrap = True
    for i, bullet in enumerate(bullets):
        para = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        para.text = bullet
        para.font.size = Pt(14)
        para.font.color.rgb = BLACK
        para.font.name = FONT_SEMIBOLD
        para.space_after = Pt(10)
        para.space_before = Pt(4)
        para.level = 0
        # Bullet formatting
        pPr = para._p.get_or_add_pPr()
        pPr.set('marL', str(228600))
        pPr.set('indent', str(-228600))
        buFont = pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'})
        buChar = pPr.makeelement(qn('a:buChar'), {'char': '\u2022'})
        for existing in pPr.findall(qn('a:buFont')):
            pPr.remove(existing)
        for existing in pPr.findall(qn('a:buChar')):
            pPr.remove(existing)
        for existing in pPr.findall(qn('a:buNone')):
            pPr.remove(existing)
        pPr.append(buFont)
        pPr.append(buChar)

    # Footer
    add_footer(slide)

    if notes:
        add_notes(slide, notes)
    return slide

def hide_slide(slide):
    slide._element.set('show', '0')

# ============================================================
# SLIDE CONTENT (preserved from original deck)
# ============================================================

# === TITLE SLIDE ===
make_title_slide(
    'Agent 365',
    'Proof of Concept',
    'Govern AI Agents with the Controls You Already Trust',
    '[Customer Name]  \u00b7  [Date]  \u00b7  4\u20136 Week Engagement',
    notes='Welcome everyone. Today we\'re kicking off the Agent 365 Proof of Concept \u2014 a focused engagement designed to show you how Microsoft brings AI agents under the same governance controls you already rely on for your people. Over the next 4 to 6 weeks, we\'ll prove this with your real agents, in your real environment.'
)

# === SLIDE 2 ===
content_slide('AI Agents Are Growing \u2014 Governance Isn\'t Keeping Up',
    ['Organizations are deploying 10\u201350+ AI agents across departments \u2014 often without IT visibility',
     'Shadow agents create blind spots in identity, data protection, and compliance',
     'Traditional tools weren\'t built for non-human actors operating at machine speed',
     'The gap between agent adoption and agent governance widens every quarter'],
    'Here\'s the challenge: your teams are already building and adopting AI agents. But most governance tooling was designed for human users. Agents don\'t fit that mold natively. The result is shadow agents, ungoverned data access, and compliance gaps that auditors will eventually find. This is the problem Agent 365 was built to solve.')

# === SLIDE 3 ===
content_slide('One Platform. Same Controls. Now for Agents.',
    ['Discover and register every agent \u2014 eliminate shadow AI',
     'Assign identity and least-privilege access through Entra ID',
     'Apply the same DLP, sensitivity labels, and insider risk policies',
     'Detect threats and respond in Defender \u2014 agents as first-class assets',
     'Maintain full audit trail using M365 compliance fabric'],
    'Agent 365 isn\'t a bolt-on product or a separate console. It extends the Microsoft 365 platform you already operate \u2014 Entra for identity, Purview for data protection, Defender for threat detection, and the M365 Admin Center for management. Your agents get the same controls as your employees.')

# === SLIDE 4 ===
content_slide('Proof of Concept \u2014 Your Agents, Your Environment, Real Results',
    ['6 use cases covering the full agent governance lifecycle',
     'Executed against 3+ agents you identify from your environment',
     'Built on your existing Entra, Purview, and Defender configurations',
     'Measurable success criteria for every use case',
     'Clear path from POC to production rollout'],
    'This isn\'t a theoretical exercise. We\'ll work with agents your teams actually use \u2014 or plan to deploy. We\'ll configure each use case against your existing security and compliance posture, measure success against specific criteria, and at the end you\'ll have a clear decision framework for broader rollout.')

# === UC1 ===
section_divider('Use Case 1', 'Agent Inventory and Discovery')

content_slide('Do You Know Every Agent Operating in Your Tenant?',
    ['Shadow agents bypass security review and operate without IT awareness',
     'Without a registry, there\'s no way to apply policy, revoke access, or audit behavior',
     'Agent 365 provides a native registry in M365 Admin Center with auto-discovery',
     'Goal: Zero shadow agents \u2014 every agent visible, classified, and accountable'],
    'Let\'s start with the most fundamental question: do you actually know how many AI agents are operating in your environment right now? Most organizations can\'t answer that confidently. Agent 365 provides a built-in registry in the M365 Admin Center that auto-discovers agents \u2014 no connectors, no third-party CMDB sync.')

content_slide('Live Demo: Agent Registry and Auto-Discovery',
    ['Navigate to Settings \u2192 Org settings \u2192 Agent management \u2192 Agent registry',
     'Review auto-discovered agents vs. manually registered agents',
     'Classify an unregistered agent (assign owner, purpose, risk tier)',
     'Demonstrate real-time discovery of a newly deployed agent',
     'Show filtering and search across agent inventory'],
    'In this demo, we\'ll open the M365 Admin Center and walk through the agent registry. You\'ll see agents that were auto-discovered alongside those that were manually registered. We\'ll classify an unregistered agent live, then deploy a test agent and watch it appear in the registry within minutes.\n\nDetailed procedure: demo-steps.md, UC1')

content_slide('How We Measure Success for Agent Discovery',
    ['100% of known agents appear in the registry within 24 hours of deployment',
     'Auto-discovery identifies at least 1 previously unknown agent in your tenant',
     'Every registered agent has an assigned owner and risk classification',
     'Zero agents operating outside the registry by end of POC'],
    'Success here is binary and measurable. By the end of this use case, every agent in your environment should be visible in the registry. We specifically look for the "aha moment" \u2014 when auto-discovery surfaces an agent that IT didn\'t know about.')

content_slide('Unified Registry \u2014 Native, Not Bolted On',
    ['Built into M365 Admin Center \u2014 same console your admins already use daily',
     'Auto-discovery is native to the platform \u2014 no API integrations or third-party scanners',
     'Agent metadata lives in the same directory as users, groups, and devices',
     'Competing solutions require separate inventory tools with manual sync'],
    'The key differentiator here is "native." You\'re not installing a third-party scanner or building API integrations to populate a CMDB. Agent discovery is built into the same Admin Center your IT team already opens every morning.\n\nCISO Talking Points:\n- "We eliminate shadow agent risk at the platform layer \u2014 agents are registered as they\'re created, not after they\'re discovered."\n- "The registry is authoritative. It doesn\'t depend on network scanning or API polling that can miss ephemeral agents."\n- "Visibility is the prerequisite for governance. Native inventory means governance starts at creation."\n\nCIO Talking Points:\n- "No new procurement for agent discovery tooling \u2014 it\'s included in your existing M365 licensing."\n- "Your IT ops team doesn\'t learn a new console."\n- "Automatic coverage means no integration project to scope, fund, or maintain."\n\nIf asked about CMDB: CMDBs track what you tell them to track. Native registry captures agents at creation \u2014 before anyone has to remember to register them.\nIf asked about non-Microsoft agents: Third-party agents can be registered via API or admin onboarding. Microsoft-platform agents require zero effort.')

# UC1-5 HIDDEN
s = content_slide('Under the Hood: Agent Discovery Flow',
    ['Agents register via Entra app registration or are detected through M365 API activity',
     'Discovery engine monitors authentication events, Graph API calls, and resource access',
     'Newly detected agents surface in registry with "Unclassified" status',
     'Admins receive alerts for unregistered agent activity',
     'Full audit trail from first detection through classification'],
    'Here\'s how it works technically. The discovery engine monitors your tenant\'s authentication and API activity. When it detects a non-human identity behaving as an agent, it surfaces that entity in the registry. Your admins get notified, classify it, and from that point forward it\'s under governance.')
hide_slide(s)

# === UC2 ===
section_divider('Use Case 2', 'Agent Identity and Least-Privilege Access')

content_slide('If Agents Don\'t Have Identity, They Can\'t Have Boundaries',
    ['Agents today often share service accounts or use over-privileged app registrations',
     'Without dedicated identity, you can\'t apply conditional access or audit individually',
     'Agent 365 gives each agent an Entra Agent ID \u2014 same identity fabric as employees',
     'Enables scoped tokens, conditional access policies, and per-agent sign-in logs'],
    'Think about how your agents authenticate today. Shared service accounts? Broad app registrations with more permissions than needed? Agent 365 gives each agent its own Entra identity \u2014 same fabric, same policies, same logs as your human users.')

content_slide('Live Demo: Entra Agent ID and Conditional Access',
    ['Create an Entra Agent ID for a test agent',
     'Assign scoped permissions using least-privilege principles',
     'Apply a conditional access policy (e.g., restrict to specific network or resource)',
     'Trigger agent activity and observe sign-in log entry',
     'Demonstrate token scoping and permission boundary enforcement'],
    'In this demo, we\'ll create a dedicated Entra Agent ID, assign it the minimum permissions needed for its function, and apply a conditional access policy. Then we\'ll trigger the agent to act, and you\'ll see its activity appear in the sign-in logs.\n\nDetailed procedure: demo-steps.md, UC2')

content_slide('How We Measure Success for Agent Identity',
    ['Each POC agent has a dedicated Entra Agent ID (no shared accounts)',
     'Conditional access policy correctly blocks agent activity outside defined boundaries',
     'Sign-in logs capture every agent authentication event with full detail',
     'Token scoping prevents access beyond assigned permissions',
     'Least-privilege verified: agent cannot access resources outside its scope'],
    'We\'ll validate that every agent in the POC has its own identity, that conditional access actually stops the agent when boundaries are crossed, and that sign-in logs capture full authentication details.')

content_slide('Why This Matters: No Parallel Identity Store',
    ['Agents live in the same Entra ID directory as your employees',
     'Same conditional access engine \u2014 no separate policy framework to maintain',
     'Sign-in logs are unified \u2014 human and agent activity in one view',
     'No third-party identity broker or agent-specific IAM tool required',
     'Your identity team already knows how to operate this'],
    'Here\'s what separates this from every other approach: there\'s no parallel identity store. Your agents are in the same directory as your people.\n\nCISO Talking Points:\n- "Agents get identities in the same Entra ID tenant your workforce uses. Conditional access, PIM, and access reviews apply without modification."\n- "Least-privilege for agents isn\'t a new initiative \u2014 it\'s an extension of the initiative you\'re already running for human identities."\n- "Workload identities for agents support the same zero-trust principles."\n\nCIO Talking Points:\n- "Zero new identity infrastructure to procure, deploy, or staff."\n- "Your IAM team\'s existing skills apply directly. No retraining, no new vendor certification."\n- "Access reviews for agents use the same workflow your managers already complete."\n\nIf asked "Agents aren\'t users": They\'re workload identities \u2014 a purpose-built identity type in Entra ID for non-human entities.\nIf asked about high-privilege agents: Same answer as for human admins: PIM with just-in-time elevation, time-bound access, and approval workflows.')

# === UC3 ===
section_divider('Use Case 3', 'Sensitive Data Protection')

content_slide('Your DLP Policies Protect Users \u2014 Do They Protect Against Agents?',
    ['AI agents process, summarize, and move sensitive data at scale',
     'Without DLP enforcement, agents can exfiltrate data faster than any insider',
     'Agent 365 extends Purview DLP and sensitivity labels to agent activities',
     'Same policies, same audit trail, same eDiscovery \u2014 now covering agents too'],
    'Your Purview DLP policies already protect against users sharing sensitive data inappropriately. But what about your agents? They process documents, summarize emails, move data between systems \u2014 all at machine speed. Agent 365 extends the same Purview policies to agent activities.')

content_slide('Live Demo: Purview DLP and Sensitivity Labels with Agents',
    ['Configure a DLP policy that applies to agent activities',
     'Trigger an agent to process a sensitivity-labeled document',
     'Observe DLP enforcement: agent action blocked or flagged',
     'Review the audit log entry showing the policy match',
     'Show insider risk signal generated by agent data handling pattern'],
    'We\'ll configure a DLP policy that explicitly covers agent activities, then ask an agent to process a document marked "Confidential." You\'ll see the policy fire \u2014 blocking the agent from sharing that content externally.\n\nDetailed procedure: demo-steps.md, UC3')

content_slide('How We Measure Success for Data Protection',
    ['DLP policy correctly blocks agent from violating data handling rules',
     'Sensitivity labels are respected by agents \u2014 labeled content is not downgraded or stripped',
     'Audit log captures every agent interaction with sensitive content',
     'Insider risk signal fires when agent exhibits anomalous data access pattern',
     'eDiscovery can surface agent-touched content using same search tools'],
    'We measure three things: enforcement (policy fires correctly), visibility (audit captures the event), and integration (same search and export tools work for agent-touched content).')

content_slide('Why This Matters: One Policy Set for Users and Agents',
    ['Same Purview DLP rules apply to both human users and AI agents',
     'Sensitivity labels travel with content regardless of who touches it',
     'eDiscovery searches naturally include agent activity \u2014 no separate export',
     'Competing solutions require duplicate policy sets with manual reconciliation'],
    'The differentiator is unification. You don\'t write one set of DLP rules for users and another for agents.\n\nCISO Talking Points:\n- "Your DLP investment extends to agents automatically. The rules you spent years tuning protect agent interactions without rework."\n- "Sensitivity labels travel with the data \u2014 whether a human or an agent touches it."\n- "Policy violations by agents produce the same alerts, in the same console, investigated by the same team."\n\nCIO Talking Points:\n- "No new data protection platform to procure. Your existing Purview investment covers agents."\n- "No \'agent DLP project\' to fund. Existing policies extend with configuration, not reimplementation."\n- "Cost of data protection governance doesn\'t scale linearly with agent count."\n\nIf asked about human-pattern DLP: Agents interact through the same channels \u2014 email, SharePoint, Teams, OneDrive. Enforcement points are identical.\nIf asked about different thresholds: Purview supports agent-scoped policy variants with different thresholds while reusing same sensitive information types.')

# === UC4 ===
section_divider('Use Case 4', 'Threat Detection and Response')

content_slide('Your SOC Monitors Users and Endpoints \u2014 Who\'s Monitoring Agents?',
    ['Compromised agents can be weaponized for lateral movement, data theft, or privilege escalation',
     'Agent-specific attack patterns (prompt injection, token theft, scope creep) need dedicated detection',
     'Agent 365 makes agents first-class Defender XDR assets with runtime protection',
     'Same SOC console, same alert queue, same incident response playbooks'],
    'Your SOC monitors users and endpoints around the clock. But what about agents? A compromised agent with broad permissions is an attacker\'s dream. Agent 365 brings agents into Defender XDR as first-class assets.')

content_slide('Live Demo: Threat Detection and Attack-Path Visualization',
    ['View agent assets in Defender XDR console',
     'Trigger a simulated anomalous agent behavior',
     'Observe alert generation in real time',
     'Walk through attack-path visualization showing potential blast radius',
     'Demonstrate SOC response: isolate agent, revoke tokens, investigate timeline'],
    'In this demo, we\'ll show your agents appearing as managed assets in the Defender console. Then we\'ll simulate anomalous behavior \u2014 an agent accessing resources outside its normal pattern. You\'ll see the alert fire, the attack-path visualization, and the SOC response workflow.\n\nDetailed procedure: demo-steps.md, UC4')

content_slide('How We Measure Success for Threat Detection',
    ['Defender generates alert within defined SLA for anomalous agent behavior',
     'Attack-path visualization accurately shows agent\'s potential blast radius',
     'SOC can isolate a compromised agent and revoke tokens from the same console',
     'Runtime protection prevents agent from executing blocked actions',
     'Incident timeline includes full agent activity history for forensics'],
    'We\'ll measure detection time, accuracy of the attack-path visualization, and the SOC\'s ability to respond from a single console. The key validation is end-to-end: anomalous behavior \u2192 alert \u2192 visualization \u2192 response \u2192 containment.')

content_slide('Why This Matters: One Console, One Incident Queue',
    ['Agents appear alongside users, devices, and apps in Defender XDR',
     'No separate SIEM integration or custom detection rules needed',
     'Attack-path analysis natively understands agent permissions and access patterns',
     'SOC response actions (isolate, revoke, block) work identically for agents',
     'Competing solutions require separate agent monitoring tools with manual SIEM forwarding'],
    'The differentiator is "first-class." Agents aren\'t bolted onto Defender as a custom data source \u2014 they\'re native assets.\n\nCISO Talking Points:\n- "Agent threats appear in the same Defender incident queue as user and endpoint threats \u2014 one priority list, one triage workflow."\n- "Correlation between agent and human activity happens automatically."\n- "Your SOC doesn\'t need an \'agent security\' specialization. Existing Defender skills apply directly."\n- "Automated response playbooks can isolate or disable a compromised agent the same way they isolate an endpoint."\n\nCIO Talking Points:\n- "No new security tool to procure for agent-specific threat detection."\n- "SOC headcount doesn\'t increase proportionally with agent count."\n- "MTTD and MTTR benefit from the same automation and AI-assisted investigation."\n\nIf asked about alert fatigue: Correlation reduces total alert volume. A compromised agent + DLP flag = one correlated incident, not three separate alerts.\nIf asked about agent-specific dashboards: Defender supports filtered views \u2014 agent-specific dashboards while benefiting from unified correlation.')

content_slide('Advanced Hunting: AIAgentsInfo Table',
    ['Defender XDR advanced hunting includes the AIAgentsInfo table \u2014 populated via Agent 365 connectors',
     'Filter with RegistrySource == "A365" for Agent 365 data',
     'Key columns: AIAgentId, AIAgentName, AgentStatus, Instructions, AgentActionTriggers, IsBlocked',
     'Enables proactive threat hunting directly from the SOC console'],
    'Advanced hunting gives your SOC team the ability to proactively query agent data alongside all other Defender XDR telemetry. The AIAgentsInfo table surfaces your full Agent 365 inventory \u2014 enabling threat hunters to identify misconfigurations, risky agents, and security gaps before they become incidents.\n\nDetailed procedure: demo-steps.md, UC4 Step 10')

content_slide('Hunting Queries: Agent Inventory and Security Risks',
    ['Query 1 \u2014 List All Agents: Full inventory with owner/creator UPNs from the SOC console',
     'Query 2 \u2014 Agents Without Instructions: Published agents missing system prompts are vulnerable to prompt injection',
     'Query 3 \u2014 MCP Tools Configured: Agents with remote MCP servers extend capabilities but increase attack surface',
     'Query 4 \u2014 Non-HTTPS Endpoints: Agents communicating over unencrypted HTTP expose data in transit'],
    'Four key hunting queries that demonstrate proactive security posture management:\n\n1. List All Agents \u2014 Visibility into the full agent inventory from within the SOC console.\n2. Agents Without Instructions \u2014 High-risk: vulnerable to prompt injection without defined behavioral boundaries. Recommendation: ensure all agents have well-defined instructions.\n3. MCP Tools Configured \u2014 MCP extends capabilities but introduces security considerations. Recommendation: confirm tools are required, review for least privilege, remove unnecessary tools.\n4. Non-HTTPS Endpoints \u2014 Exposes data to interception/tampering. Recommendation: update all HTTP actions to HTTPS.\n\nThese queries shift the SOC from reactive alerting to proactive posture management for AI agents.')

# === UC5 ===
section_divider('Use Case 5', 'Audit, Compliance, and Lifecycle')

content_slide('When Auditors Come Knocking, Can You Show the Agent Evidence Chain?',
    ['Regulatory frameworks increasingly expect governance over autonomous AI systems',
     'Auditors need a complete trail: who deployed it, what it accessed, when it was deactivated',
     'Agent 365 provides full lifecycle audit through M365 compliance fabric',
     'Same evidence format auditors already accept for user activities'],
    'Here\'s a scenario playing out right now: your compliance team passes an audit, but the auditor asks, "What about your AI agents?" If you can\'t answer that in the same compliance framework they already trust, you have a finding.')

content_slide('Live Demo: Full Audit Trail and Lifecycle Management',
    ['Review complete audit trail for an agent (deployment through operation)',
     'Execute lifecycle drill: block agent, deactivate agent, delete agent',
     'Verify each lifecycle action generates audit entry with timestamp and actor',
     'Show data retention policy applying to agent audit records',
     'Demonstrate compliance report including agent activities'],
    'We\'ll walk through the full lifecycle of an agent in audit terms. You\'ll see the trail from deployment through every action it took. Then we\'ll execute a lifecycle drill \u2014 blocking, deactivating, and deleting \u2014 verifying each action is audited.\n\nDetailed procedure: demo-steps.md, UC5')

content_slide('How We Measure Success for Audit and Compliance',
    ['Complete audit trail exists from agent deployment through deletion',
     'Each lifecycle state change generates a timestamped audit entry',
     'Data retention policies correctly apply to agent audit records',
     'Compliance report export includes agent activities in accepted format',
     'No gaps in evidence chain \u2014 auditor can reconstruct full agent history'],
    'Success here means an auditor can pick up the compliance report, see agent activities alongside user activities, and trace any agent\'s full history from deployment to deletion.')

content_slide('Why This Matters: No Parallel Evidence Chain',
    ['Agent audit data feeds the same M365 compliance fabric used for users',
     'Auditors already accept this format \u2014 no new validation required',
     'Data retention, legal hold, and eDiscovery apply uniformly',
     'No separate compliance tool generating evidence that auditors must learn to trust',
     'Competing solutions create a second evidence chain requiring separate attestation'],
    'The power here is trust. Your auditors already accept evidence from M365 compliance tools. Agent activities feed that same system.\n\nCISO Talking Points:\n- "Agent lifecycle events are captured in the same unified audit log your compliance team already queries."\n- "Retention policies that apply to human activity audit data apply to agent activity audit data."\n- "When auditors ask \'how do you govern agents?\', the answer is \'the same way we govern everything else.\'"\n\nCIO Talking Points:\n- "No new compliance tooling procurement. Agent audit evidence lives in existing M365 compliance infrastructure."\n- "Your compliance team doesn\'t need new skills or certifications."\n- "Audit preparation cost doesn\'t increase with agent adoption."\n\nIf asked about auditor unfamiliarity: They\'ll question it less if it\'s in a system they already trust.\nIf asked about agent-specific reports: The unified audit log supports filtered queries for agent-specific views.')

# UC5-5 HIDDEN
s = content_slide('Under the Hood: How Agent Lifecycle Feeds Compliance',
    ['Every agent action generates a unified audit log entry (same schema as user events)',
     'Lifecycle state changes trigger compliance workflow notifications',
     'Retention labels auto-apply to agent audit records based on classification',
     'eDiscovery index includes agent-generated content and interactions',
     'Export APIs produce audit packages in standard regulatory formats'],
    'Technically, every agent action writes to the same unified audit log as user actions \u2014 same schema, same retention handling. When an agent changes lifecycle state, compliance workflows fire automatically.')
hide_slide(s)

# === UC6 ===
section_divider('Use Case 6', 'Productivity Integration')

content_slide('Agents in Outlook, Teams, Word, and SharePoint \u2014 Governed or Ungoverned?',
    ['AI agents increasingly operate within productivity tools \u2014 drafting emails, posting in Teams, editing documents',
     'Without governance, these agents act with unbounded access across your collaboration surface',
     'Agent 365 enables governed connections with real-time allow/block enforcement',
     'Same productivity surface users know \u2014 no separate agent UI or plug-in'],
    'Your users interact with AI agents inside the apps they use every day \u2014 Outlook, Teams, Word, SharePoint. But are those agent interactions governed? Agent 365 enables real-time allow/block enforcement within the productivity surface.')

content_slide('Live Demo: Governed Agent Actions in M365 Apps',
    ['Deploy a test agent with connections to Teams and SharePoint',
     'Configure allow/block policies for specific agent actions',
     'Trigger an allowed action \u2014 agent posts summary to a Teams channel',
     'Trigger a blocked action \u2014 agent attempts to share externally',
     'Observe real-time enforcement and audit trail for both actions'],
    'We\'ll deploy a test agent connected to Teams and SharePoint, configure policies, and demonstrate both allowed and blocked actions in real time.\n\nDetailed procedure: demo-steps.md, UC6')

content_slide('How We Measure Success for Productivity Integration',
    ['Agent successfully operates within allowed boundaries in Outlook, Teams, Word, and/or SharePoint',
     'Blocked actions are denied in real time (sub-second enforcement)',
     'Allow/block policies can be changed and take effect without redeploying the agent',
     'Audit trail captures both allowed and blocked actions with full context',
     'End users experience no disruption \u2014 governance is invisible to them'],
    'We need to prove three things: allowed actions work smoothly, blocked actions are denied instantly, and policy changes take effect without requiring agent redeployment.')

content_slide('Why This Matters: No Separate Agent UI',
    ['Agents operate natively within M365 apps \u2014 no separate interface for users to learn',
     'Governance is platform-level \u2014 not a per-app plug-in that can be bypassed',
     'Allow/block enforcement is real-time and centrally managed',
     'Users interact with agent output in the apps they already use daily',
     'Competing solutions require browser extensions, sidebars, or standalone agent portals'],
    'The differentiator is seamlessness. Your users don\'t need a new app, a browser extension, or a sidebar.\n\nCISO Talking Points:\n- "Agents in familiar surfaces means users follow established security behaviors."\n- "No new client application means no new attack surface to secure, patch, or monitor."\n- "Security controls that apply to Teams and M365 apps apply to agent interactions within them."\n\nCIO Talking Points:\n- "Adoption cost is near zero. Users interact with agents in Teams and M365 apps they already use daily."\n- "No client deployment project. No MDM profile updates."\n- "IT admin overhead is configuration, not infrastructure."\n\nIf asked about specialized interfaces: Adaptive Cards in Teams support complex workflows without leaving the familiar surface.\nIf asked about controlling visibility: Agent visibility controlled through same Teams app permission policies and Entra ID group assignments.')

# === CLOSING ===
section_divider('Summary & Next Steps', 'From Proof of Concept to Production')

content_slide('What We Proved: Agent Governance That Works Today',
    ['\u2705 UC1: Every agent discovered and registered \u2014 zero shadow agents',
     '\u2705 UC2: Dedicated identity with least-privilege \u2014 same fabric as employees',
     '\u2705 UC3: DLP and sensitivity labels enforced \u2014 same policies as users',
     '\u2705 UC4: Threat detection and response \u2014 agents as first-class Defender assets',
     '\u2705 UC5: Full audit trail \u2014 compliance fabric auditors already trust',
     '\u2705 UC6: Governed productivity integration \u2014 real-time allow/block, no plug-ins'],
    'Let\'s bring it all together. Over these weeks, we proved that Agent 365 delivers real, measurable governance across every dimension. And we did it using the tools and frameworks your teams already operate. No new consoles, no parallel systems, no learning curve.')

content_slide('Broader Rollout: Extending Agent Governance Across Your Organization',
    ['Expand from 3+ POC agents to full agent estate',
     'Integrate agent governance into existing change management processes',
     'Align agent lifecycle policies with your compliance calendar',
     'Enable self-service agent registration with guardrails for development teams',
     'Establish ongoing monitoring cadence with SOC team'],
    'The POC proved the technology works with your agents in your environment. The next step is broader rollout \u2014 extending governance to your full agent estate, integrating with change management, and enabling self-service registration.')

content_slide('Let\'s Keep the Conversation Going',
    ['Your Microsoft account team: [Account Team Contacts]',
     'Agent 365 documentation: [Link to docs]',
     'POC completion report: [Will be delivered at engagement close]',
     'FastTrack support for production deployment: [Link/contact]',
     'Community and feedback: [Tech Community link]'],
    'Thank you for your time and partnership throughout this Proof of Concept. Your account team is here to support next steps. We\'ll deliver the formal POC completion report within one week of engagement close. Let\'s govern your agents the way you govern your people.')

# === APPENDIX ===
section_divider('Appendix', 'Reference Materials')

content_slide('Roles & Responsibilities',
    ['Executive Sponsor \u2014 Approves scope, provides access, makes go/no-go decision',
     'IT Admin \u2014 Configures M365 Admin Center, manages agent registry',
     'Identity Admin \u2014 Manages Entra ID agent identities, conditional access policies',
     'Data Security Lead \u2014 Configures Purview DLP, sensitivity labels for agents',
     'SOC Analyst \u2014 Monitors Defender alerts, investigates agent threats, runs response',
     'Business Owners \u2014 Identify agents for POC, validate success criteria alignment'],
    'Reference slide: Roles involved in the Agent 365 Proof of Concept engagement.')

content_slide('Agent Types in Scope',
    ['Microsoft 365 Copilot',
     'Copilot Agent Builder agents',
     'Copilot Studio agents',
     'Microsoft 365 Copilot Chat',
     'SharePoint agents',
     'Azure AI Foundry agents',
     'First-party Microsoft agents'],
    'Reference slide: Agent types supported in the Agent 365 POC. Note: Agent 365 is the governance capability set within Microsoft 365 \u2014 not a separately licensed product.')

content_slide('Glossary',
    ['Agent 365 \u2014 Microsoft 365 governance capability set for AI agents (not a separate SKU)',
     'Entra Agent ID \u2014 Workload identity in Entra ID purpose-built for agent governance',
     'Unified Audit Log \u2014 M365 compliance system of record for all user and agent activity',
     'Conditional Access \u2014 Policy engine enforcing identity-based access rules',
     'DLP (Data Loss Prevention) \u2014 Purview policies preventing unauthorized data exposure',
     'Sensitivity Labels \u2014 Classification and protection tags that travel with content'],
    'Reference slide: Key terms for mixed-audience understanding.')

# Save
output_path = r'C:\Users\chrwilliams\Agent-365-POC\Agent365-POC-Deck.pptx'
prs.save(output_path)
print(f'SUCCESS: Saved to {output_path}')
print(f'Total slides: {len(prs.slides)}')
